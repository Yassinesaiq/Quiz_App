from django.contrib.auth.models import Group, User
from rest_framework import viewsets ,status
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from .serializers import *
from .models import *
from django.shortcuts import *
from .forms import *
from django.db import transaction
from django.contrib.auth.decorators import login_required ,user_passes_test
import logging
from django.contrib.auth import logout , authenticate , login
from django.contrib import messages
from django.views.decorators.http import require_POST
from django.utils import timezone
from datetime import timedelta
from rest_framework.authtoken.models import Token
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated ,AllowAny
from rest_framework.authentication import TokenAuthentication
from openpyxl import Workbook
from django.utils.timezone import localtime
from rest_framework.decorators import api_view, permission_classes
from django.contrib.auth.models import User
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill
from django.contrib.admin.views.decorators import staff_member_required
from rest_framework import generics
from .models import TextQuestion
from .serializers import TextQuestionSerializer
from rest_framework.response import Response
from rest_framework import status
from django.urls import reverse
from django.forms import modelform_factory, inlineformset_factory
from django.contrib.admin.views.decorators import staff_member_required




class UserViewSet(viewsets.ModelViewSet):
    """
    API endpoint that allows users to be viewed or edited.
    """
    queryset = User.objects.all().order_by('-date_joined')
    serializer_class = UserSerializer
    permission_classes = [IsAuthenticated]


class GroupViewSet(viewsets.ModelViewSet):
    """
    API endpoint that allows groups to be viewed or edited.
    """
    queryset = Group.objects.all().order_by('name')
    serializer_class = GroupSerializer
    permission_classes = [IsAuthenticated]


class CustomPagination(PageNumberPagination):
    def get_paginated_response(self, data):
        return Response(data)
    
class QuestionsViewset(viewsets.ModelViewSet):
    queryset = Questions.objects.all()
    serializer_class = QuestionSerializer
    pagination_class = CustomPagination

    def get_queryset(self):
        queryset = super().get_queryset()
        topic_id = self.request.query_params.get('topic_id')
        if topic_id is not None:
            queryset = queryset.filter(topic_id=topic_id)
        return queryset

class TopicViewset(viewsets.ModelViewSet):
    queryset = Topics.objects.filter(visible=True) 
    serializer_class = TopicSerializer
    pagination_class = CustomPagination
    permission_classes = [IsAuthenticated]

class TextQuestionViewset(viewsets.ModelViewSet):
    queryset = TextQuestion.objects.all()
    serializer_class = TextQuestionSerializer
    pagination_class = CustomPagination
    permission_classes = [IsAuthenticated]

@login_required   
@staff_member_required(login_url='/login/')
def topics_list(request):
    topic_type = request.GET.get('topic_type')
    if topic_type:
        topics = Topics.objects.filter(topic_type__name=topic_type)
    else:
        topics = Topics.objects.all()
    return render(request, 'topics_list.html', {'topics': topics, 'topic_type': topic_type})

@login_required   
@staff_member_required(login_url='/login/')
def topic_questions(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    questions = Questions.objects.filter(topic_id=topic)
    topic_type = topic.topic_type.name if topic.topic_type else "MCQ"
    if topic_type == "Textfrage":
        return redirect("view_text_questions_of_topic", topic_id=topic.pk)
    elif topic_type == "MCQ + Text":
        # Replace 'mcq_text_question_view' with your actual view name for MCQ + Text
        return redirect("view_mcq_text_questions", topic_id=topic.pk)
    return render(request, 'topic_questions.html', {
        'topic': topic,
        'questions': questions,
        'topic_type': topic_type
    })

# ---------------- Topics ---------------- #
@login_required   
@staff_member_required(login_url='/login/')
def add_topic(request):
    #  topic_type VORHER initialisieren, damit es immer existiert
    topic_type = None
    type_id = request.GET.get("type_id")
    if type_id:
        try:
            topic_type = TopicType.objects.get(id=type_id)
        except TopicType.DoesNotExist:
            topic_type = None

    if request.method == "POST":
        form = TopicForm(request.POST, request.FILES)
        if form.is_valid():
            topic = form.save(commit=False)
            topic.created_by = request.user
            topic.created_at = timezone.now() + timedelta(hours=2)
            if topic_type:
                topic.topic_type = topic_type  #Type speichern
            topic.save()
            form.save()
            return redirect('topics_list')
    else:
        form = TopicForm()

    return render(
        request,
        'add_edit_topic.html',
        {
            'form': form,
            'title': 'Thema hinzufügen',
            'topic_type': topic_type,  # immer gesetzt (None wenn nicht gefunden)
        }
    )

@login_required   
@staff_member_required(login_url='/login/')
def edit_topic(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    topic_type = topic.topic_type  #  Typ des Themas beibehalten

    if request.method == "POST":
        form = TopicForm(request.POST, request.FILES, instance=topic)
        if form.is_valid():
            topic = form.save(commit=False)
            topic.created_by = request.user
            topic.created_at = timezone.now() + timedelta(hours=2)
            topic.topic_type = topic_type  #  Sicherstellen, dass Typ nicht verloren geht
            topic.save()
            return redirect('topics_list')
    else:
        form = TopicForm(instance=topic)

    return render(
        request,
        'add_edit_topic.html',
        {
            'form': form,
            'title': 'Thema bearbeiten',
            'topic_type': topic_type  #  Im Template verfügbar
        }
    )

@login_required   
@staff_member_required(login_url='/login/')
def delete_topic(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    with transaction.atomic():
        #  Alle MCQ-Fragen löschen
        Questions.objects.filter(topic_id=topic).delete()

        #  Alle Textfragen löschen (falls vorhanden)
        TextQuestion.objects.filter(topic=topic).delete()

        #  Danach das Thema selbst löschen
        topic.delete()

    return redirect('topics_list')


# ---------------- Questions ---------------- #
@login_required   
@staff_member_required(login_url='/login/')
def add_question(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    topic_type = topic.topic_type.name if topic.topic_type else "MCQ"
    if topic_type == "Textfrage":
        return redirect("add_text_questions_to_topic", topic_id=topic.pk)
    if topic_type == "MCQ + Text":
        return redirect("add_mcq_text_question_view", topic_id=topic.pk)

    #  Standard: MCQ-Formular
    if request.method == "POST":
        form = QuestionForm(request.POST, request.FILES, current_topic=topic)
        if form.is_valid():
            question = form.save(commit=False)
            question.topic_id = topic
            question.save()

            if "save_continue" in request.POST:
                return redirect("add_question", topic_id=topic.pk)  # weitere MCQ-Frage
            else:
                return redirect("topic_questions", topic_id=topic.pk)  # zurück zur Übersicht
    else:
        form = QuestionForm(current_topic=topic)

    return render(
        request,
        "add_edit_question.html",
        {
            "form": form,
            "title": f"MCQ-Frage hinzufügen zu: {topic.topic}",
            "topic": topic,
            "topic_type": topic_type,
        },
    )


@login_required   
@staff_member_required(login_url='/login/')
def edit_question(request, question_id):
    question = get_object_or_404(Questions, pk=question_id)

    if request.method == "POST":
        form = QuestionForm(request.POST, request.FILES, instance=question)
        # Ensure topic_id is not changed
        form.fields["topic_id"].disabled = True
        if form.is_valid():
            # Do not update topic_id even if POST data contains it
            edited_question = form.save(commit=False)
            edited_question.topic_id = question.topic_id
            edited_question.save()
            return redirect('topic_questions', topic_id=question.topic_id.pk)
    else:
        form = QuestionForm(instance=question)
        # Disable topic_id field in the form
        if "topic_id" in form.fields:
            form.fields["topic_id"].disabled = True
        logger = logging.getLogger(__name__)
        logger.info(question.des_img)
    return render(request, 'add_edit_question.html', {
        'form': form,
        'title': 'Frage bearbeiten',
        'topic': question.topic_id
    })

@login_required   
@staff_member_required(login_url='/login/')
def delete_question(request, question_id):
    question = get_object_or_404(Questions, pk=question_id)
    topic_id = question.topic_id.topic_id
    question.delete()
    return redirect('topic_questions', topic_id=topic_id)


@login_required   
@staff_member_required(login_url='/login/')
def dashboard(request):
    topics_count = Topics.objects.count()
    questions_count = Questions.objects.count()
    latest_topics = Topics.objects.filter(created_by=request.user).order_by('-created_at')[:5]

    return render(request, 'dashboard.html', {
        'topics_count': topics_count,
        'questions_count': questions_count,
        'latest_topics': latest_topics,
    })


@require_POST
def logout_view(request):
    logout(request)  # This removes the session
    messages.success(request, "Sie wurden erfolgreich abgemeldet.")
    return redirect('login')  # or use 'home' or a custom page


@login_required   
@staff_member_required(login_url='/login/')
def toggle_topic_visibility(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    topic.visible = not topic.visible
    topic.save()
    return redirect('topics_list')


from django.contrib.auth import authenticate, login
from django.http import HttpResponse
from django.shortcuts import render, redirect


def login_user(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')

        print(f"DEBUG → username: {username}, password: {password}")  # optional debug

        user = authenticate(request, username=username, password=password)

        if user is not None:
            login(request, user)
            groups = [g.name.lower() for g in user.groups.all()]
            print("DEBUG → groups:", groups)

            # ✅ redirect by role
            if "azubis" in groups:
                return redirect('azubi_dashboard')
            elif user.is_superuser or user.is_staff:
                return redirect('dashboard')
            else:
                return redirect('user_profile')
        else:
            messages.error(request, "Ungültiger Benutzername oder Passwort.")
            print("DEBUG → Authentication failed")
            return render(request, 'login.html', {'form': {}})

    return render(request, 'login.html', {'form': {}})

    
        
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_current_user(request):
    user = request.user
    data = [{
        "username": user.username,
        "first_name": user.first_name,
        "last_name": user.last_name
        
    }]
    return Response(data, status=status.HTTP_200_OK)


class SubmitResultView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request):
        print("Authorization Header:", request.META.get('HTTP_AUTHORIZATION'))
        serializer = QuizResultSerializer(data=request.data, context={'request': request})
        if serializer.is_valid():
            serializer.save()
            return Response({'status': 'success'}, status=201)
        return Response(serializer.errors, status=400)

from django.db.models import Q
@login_required   
@staff_member_required(login_url='/login/')
def score_list_view(request):
    # --- Basis-Query ---
    sessions = (
        TestSession.objects
        .select_related("user", "user__profile", "topic__topic_type")  # ✅ Profil des Users mitladen
        .prefetch_related("quiz_results__wrong_answers", "text_answers__question")
        .order_by("-started_at")
    )

    # --- Filter ---
    topic_filter = request.GET.get("topic")
    type_filter = request.GET.get("topic_type")
    first_name_filter = request.GET.get("first_name")
    last_name_filter = request.GET.get("last_name")

    if topic_filter:
        sessions = sessions.filter(topic__topic__icontains=topic_filter)
    if type_filter:
        sessions = sessions.filter(topic__topic_type__name__icontains=type_filter)
    if first_name_filter:
        sessions = sessions.filter(user__first_name__icontains=first_name_filter)
    if last_name_filter:
        sessions = sessions.filter(user__last_name__icontains=last_name_filter)

    # --- Vorschläge für Datalists ---
    topic_suggestions = Topics.objects.values_list("topic", flat=True).distinct()
    type_suggestions = TopicType.objects.values_list("name", flat=True).distinct()
    first_name_suggestions = User.objects.exclude(first_name="").values_list("first_name", flat=True).distinct()
    last_name_suggestions = User.objects.exclude(last_name="").values_list("last_name", flat=True).distinct()

    # --- Ergebnisse aufbereiten ---
    results = []
    for session in sessions:
        quiz_result = session.quiz_results.first()
        text_answers = session.text_answers.all()

        # MCQ-Teil
        quiz_score = quiz_result.score if quiz_result else 0
        quiz_max = quiz_result.maxsize if quiz_result else 0

        # Text-Teil
        total_text_score = sum(a.score for a in text_answers if a.score is not None)
        total_text_max_score = sum(a.question.max_score for a in text_answers)

        # Gesamt
        total_score = quiz_score + total_text_score
        total_max = quiz_max + total_text_max_score

        results.append({
            "session_id": session.id,
            "user": session.user,
            "user_profile": getattr(session.user, "profile", None),  # ✅ sicherstellen, dass ein Profil existiert
            "topic": session.topic.topic,
            "topic_type": session.topic.topic_type.name if session.topic.topic_type else "Unbekannt",
            "quiz_result": quiz_result,
            "wrong_answers": quiz_result.wrong_answers.all() if quiz_result else [],
            "text_answers": text_answers,
            "pending_count": text_answers.filter(score__isnull=True).count(),
            "total_text_score": total_text_score,
            "total_text_max_score": total_text_max_score,
            "score": quiz_score,
            "maxsize": quiz_max,
            "total_score": total_score,
            "total_max_score": total_max,
            "created_at": localtime(session.started_at + timedelta(hours=2)),
            "deletion_date": localtime(session.started_at + timedelta(days=60)),
        })

    # --- Template Rendern ---
    return render(request, "score_list.html", {
        "results": results,
        "topic_suggestions": topic_suggestions,
        "type_suggestions": type_suggestions,
        "first_name_suggestions": first_name_suggestions,
        "last_name_suggestions": last_name_suggestions,
    })


#Excel Datei Herunterladen 
#https://openpyxl.readthedocs.io/en/stable/
@login_required   
@staff_member_required(login_url='/login/')
def download_result_excel(request, id):
    if not request.user.is_staff:
        return HttpResponse("Nicht erlaubt", status=403)

    result = get_object_or_404(QuizResult, id=id)
    session = result.session  # damit wir die dazugehörigen TextAnswers bekommen

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Testergebnis"

    bold_font = Font(bold=True)

    # === Kopfbereich ===
    sheet["A1"] = "Benutzer"
    sheet["B1"] = result.user.get_full_name() or result.user.username
    sheet["A2"] = "Thema"
    sheet["B2"] = str(result.topic)
    sheet["A3"] = "Gesamtpunkte"
    sheet["B3"] = result.score
    sheet["A4"] = "Datum"
    sheet["B4"] = result.created_at.strftime('%d.%m.%Y %H:%M')

    for cell in ["A1", "A2", "A3", "A4"]:
        sheet[cell].font = bold_font

    sheet.append([])

    # === MCQ-Falsche Antworten ===
    sheet.append(["MCQ - Frage", "Richtige Antwort", "Gegebene Antwort"])
    last_row = sheet.max_row
    for col in range(1, 4):
        sheet.cell(row=last_row, column=col).font = bold_font

    if result.wrong_answers.exists():
        for wrong in result.wrong_answers.all():
            sheet.append([wrong.question, wrong.correct_answer, wrong.selected_option])
    else:
        sheet.append(["Alle Fragen korrekt beantwortet 🎉", "", ""])

    # === Textantworten ===
    if session:
        text_answers = session.text_answers.select_related("question").all()
        if text_answers.exists():
            sheet.append([])  # Leerzeile
            sheet.append(["TEXTFRAGEN", "Bewertung", "Feedback"])
            last_row = sheet.max_row
            for col in range(1, 4):
                sheet.cell(row=last_row, column=col).font = bold_font
            for ans in text_answers:
                sheet.append([
                    ans.question.question_text,
                    ans.score if ans.score is not None else "Noch nicht bewertet",
                    ans.feedback or ""
                ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="{result.user.username}_result_{result.id}.xlsx"'

    workbook.save(response)
    return response

@login_required   
@staff_member_required(login_url='/login/')
def download_all_results_excel(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Alle Ergebnisse"

    # Formatierungen
    bold_font = Font(bold=True)
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    blue_fill = PatternFill(start_color="90D5FF", end_color="90D5FF", fill_type="solid")

    # Kopfzeilen
    headers = [
        "Benutzer", "Vorname", "Nachname",
        "Thema", "Thema-Typ", "Score", "Datum",
        "Fragetyp", "Frage", "Gegebene Antwort", "Richtige Antwort / Max Score", "Punkte / Feedback"
    ]
    sheet.append(headers)
    for col in range(1, len(headers) + 1):
        sheet.cell(row=1, column=col).font = bold_font

    # Ergebnisse
    results = QuizResult.objects.select_related("user", "session").prefetch_related("wrong_answers")

    for result in results:
        user = result.user
        topic_name = result.topic  # vermutlich CharField
        topic_obj = Topics.objects.filter(topic=topic_name).select_related("topic_type").first()
        topic_type = topic_obj.topic_type.name if topic_obj and topic_obj.topic_type else "Unbekannt"

        # 1. MCQ-Ergebnisse
        wrongs = result.wrong_answers.all()
        if wrongs.exists():
            for wrong in wrongs:
                sheet.append([
                    user.username,
                    user.first_name,
                    user.last_name,
                    topic_name,
                    topic_type,
                    result.score,
                    result.created_at.strftime("%d.%m.%Y %H:%M"),
                    "MCQ",
                    wrong.question,
                    wrong.selected_option,
                    wrong.correct_answer,
                    "Falsch beantwortet"
                ])
        else:
            sheet.append([
                user.username,
                user.first_name,
                user.last_name,
                topic_name,
                topic_type,
                result.score,
                result.created_at.strftime("%d.%m.%Y %H:%M"),
                "MCQ",
                "-",
                "-",
                "Alle korrekt",
                "✔️"
            ])

        # 2. Textfragen dieser Session
        if result.session:
            text_answers = TextAnswer.objects.filter(session=result.session).select_related("question")
            for ans in text_answers:
                sheet.append([
                    user.username,
                    user.first_name,
                    user.last_name,
                    topic_name,
                    topic_type,
                    result.score,
                    result.created_at.strftime("%d.%m.%Y %H:%M"),
                    "Textfrage",
                    ans.question.question_text if ans.question else "Unbekannt",
                    ans.answer_text,
                    f"Max: {ans.question.max_score if ans.question else '-'}",
                    f"{ans.score if ans.score is not None else '-'} | {ans.feedback or ''}"
                ])

    # Farben setzen
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        score_cell = row[5]  # Score-Spalte
        score_cell.fill = blue_fill

        qtype = row[7].value
        if qtype == "MCQ":
            gegeben_cell = row[9]
            richtig_cell = row[10]
            if row[11].value == "Falsch beantwortet":
                gegeben_cell.fill = red_fill
                richtig_cell.fill = green_fill

    # Datei zurückgeben
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="alle_ergebnisse.xlsx"'
    workbook.save(response)
    return response




@login_required   
@staff_member_required(login_url='/login/')
@transaction.atomic
def delete_result(request, id):
    session = get_object_or_404(TestSession.objects.select_related("topic__topic_type"), id=id)
    topic_type = session.topic.topic_type.name if session.topic and session.topic.topic_type else "MCQ"

    # Fall 1: Nur MCQ
    if topic_type == "MCQ":
        session.quiz_results.all().delete()

    # Fall 2: Nur Text
    elif topic_type == "Textfrage":
        session.text_answers.all().delete()

    # Fall 3: Mixed (MCQ + Text)
    elif topic_type == "MCQ + Text":
        session.quiz_results.all().delete()
        session.text_answers.all().delete()

    # Session selbst löschen
    session.delete()

    return redirect("score_list_view")



from django.db import transaction

@login_required   
@staff_member_required(login_url='/login/')
@transaction.atomic
def delete_all_result(request):
    #  Alle Sessions zuerst löschen (löscht automatisch TextAnswers, wenn on_delete=CASCADE)
    TestSession.objects.all().delete()
    # alle QuizResults löschen (löscht WrongAnswers per on_delete=CASCADE)
    QuizResult.objects.all().delete()
    return redirect("score_list_view")

@login_required   
@staff_member_required(login_url='/login/')
def text_answers_list(request):
    session_id = request.GET.get("session")
    if not session_id:
        return HttpResponse("Session nicht angegeben", status=400)

    answers = TextAnswer.objects.filter(session_id=session_id).select_related("question", "user")
    return render(request, "text_answers_list.html", {
        "answers": answers,
        "session_id": session_id
    })


def review_text_answer(request, answer_id):
    answer = get_object_or_404(TextAnswer, pk=answer_id)
    session_id = request.GET.get("session") or request.POST.get("session")

    if not session_id:
        return HttpResponse("Session nicht angegeben", status=400)

    if request.method == "POST":
        score = request.POST.get("score")
        feedback = request.POST.get("feedback")
        max_score = answer.question.max_score
        if score:
            score_val = float(score)
            answer.score = min(score_val, max_score)  # Prevent exceeding max score
        answer.feedback = feedback
        answer.is_reviewed = True
        answer.save()

        return redirect(f"{reverse('text_answers_list')}?session={session_id}")

    return render(request, "review_text_answer.html", {"answer": answer})


class TextQuestionByTopicNameView(generics.ListAPIView):
    serializer_class = TextQuestionSerializer
    pagination_class = None   #  Wichtig: keine Pagination

    def get_queryset(self):
        topic_name = self.request.query_params.get("topic")
        if topic_name:
            return TextQuestion.objects.filter(topic__topic__iexact=topic_name)
        return TextQuestion.objects.none()

    

class SubmitTextAnswerAPI(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        question_id = request.data.get("question")
        answer_text = request.data.get("answer_text")
        session_id = request.data.get("session")  # New: session ID from Android
        topic_name = request.data.get("topic")    # optional fallback to auto-create session

        if not question_id or not answer_text:
            return Response({"error": "question and answer_text are required"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            question = TextQuestion.objects.get(id=question_id)
        except TextQuestion.DoesNotExist:
            return Response({"error": "Question not found"}, status=status.HTTP_404_NOT_FOUND)

        #  Ensure session exists (create one if none provided)
        session = None
        if session_id:
            session = TestSession.objects.filter(id=session_id, user=request.user).first()

        if not session:
            if topic_name:
                try:
                    topic = Topics.objects.get(topic=topic_name)
                except Topics.DoesNotExist:
                    return Response({"error": "Topic not found"}, status=status.HTTP_404_NOT_FOUND)
                session = TestSession.objects.create(user=request.user, topic=topic)
            else:
                return Response({"error": "Session not found and no topic provided"}, status=status.HTTP_400_BAD_REQUEST)

        # Speichern der Textantwort
        text_answer = TextAnswer.objects.create(
            user=request.user,
            question=question,
            answer_text=answer_text,
            session=session  #  Link to session
        )

        return Response(
            {"success": "Answer saved", "session_id": session.id},  #  Send back session ID
            status=status.HTTP_201_CREATED
        )



@api_view(["POST"])
@permission_classes([IsAuthenticated])
def start_test_session(request):
    topic_name = request.data.get("topic")
    topic = Topics.objects.filter(topic=topic_name).first()
    if not topic:
        return Response({"error": "Topic not found"}, status=404)
    
    """  prüfen ob der User schon eine Session für dieses Topic hat
    existing_session = TestSession.objects.filter(user=request.user, topic=topic).first()
    if existing_session:
        return Response({
            "error": "Du hast diesen Test bereits gemacht."
        }, status=403) """

    session = TestSession.objects.create(user=request.user, topic=topic)
    return Response({ "message": "Session gestartet",
                      "session_id": session.id})


@login_required   
@staff_member_required(login_url='/login/')
def delete_session(request, session_id):
    session = get_object_or_404(TestSession, pk=session_id)

    # Optional: Nur Admins oder Besitzer dürfen löschen
    if not request.user.is_staff and session.user != request.user:
        return HttpResponse("Nicht erlaubt", status=403)

    session.delete()
    return redirect('score_list_view')


@login_required   
@staff_member_required(login_url='/login/')
def choose_topic_type(request):
    if request.method == "POST":
        choice = request.POST.get("choice")
        try:
            topic_type = TopicType.objects.get(name__iexact=choice)  # Hole den Typ aus DB
        except TopicType.DoesNotExist:
            topic_type = None

        if topic_type:
            # Weiterleiten mit Typ-ID
            return redirect(f"{reverse('add_topic')}?type_id={topic_type.id}")

    return render(request, "choose_topic_type.html")
    

@login_required   
@staff_member_required(login_url='/login/')
def add_text_question(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    topic_type = topic.topic_type.name if topic.topic_type else "MCQ"

    # Only allow adding text questions for correct topic types
    if topic_type not in ["Textfrage", "MCQ + Text"]:
        messages.error(request, "Für diesen Thema-Typ können keine Textfragen hinzugefügt werden.")
        return redirect("topics_list")

    if request.method == "POST":
        form = TextQuestionForm(request.POST, request.FILES)
        if form.is_valid():
            question = form.save(commit=False)
            question.topic = topic
            question.save()

            if "save_continue" in request.POST:
             return redirect("add_text_question", topic_id=topic.topic_id)
            else:
                messages.success(request, "Frage gespeichert.")
                if topic_type == "MCQ + Text":
                    return redirect("view_mcq_text_questions", topic_id=topic.topic_id)
                else:
                 return redirect("view_text_questions_of_topic", topic_id=topic.topic_id)
        else:
            messages.error(request, "Bitte korrigiere die Fehler im Formular.")
    else:
        form = TextQuestionForm()

    return render(request, "add_text_question.html", {
        "form": form,
        "topic": topic,
        "topic_type": topic_type,
    })



@login_required   
@staff_member_required(login_url='/login/')
def view_text_questions_of_topic(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)
    text_questions = TextQuestion.objects.filter(topic=topic)
    return render(request, "view_text_questions.html", {
        "topic": topic,
        "text_questions": text_questions
    })

@login_required   
@staff_member_required(login_url='/login/')
def edit_text_question(request, question_id):
    question = get_object_or_404(TextQuestion, pk=question_id)
    topic_type = question.topic.topic_type
    if request.method == "POST":
        form = TextQuestionForm(request.POST, request.FILES, instance=question)
        if form.is_valid() and topic_type == "Textfrage":
            form.save()
            messages.success(request, "Textfrage wurde bearbeitet.")
            return redirect("view_text_questions_of_topic", topic_id=question.topic.pk)
        elif form.is_valid() and topic_type == "MCQ + Text":
             messages.success(request, "Textfrage wurde bearbeitet.")
             return redirect("view_mcq_text_questions", topic_id=question.topic.pk)
        else :
            messages.error(request, "Bitte korrigiere die Fehler.")
    else:
        form = TextQuestionForm(instance=question)
    return render(request, "edit_text_question.html", {
        "form": form,
        "question": question,
        "topic": question.topic
    })


@login_required   
@staff_member_required(login_url='/login/')
def delete_text_question(request, question_id):
    question = get_object_or_404(TextQuestion, pk=question_id)
    tp_type = question.topic.topic_type.name
    topic_id = question.topic.pk
    question.delete()
    messages.success(request, "Textfrage wurde gelöscht.")
    if tp_type == "MCQ + Text":
        return redirect("view_mcq_text_questions", topic_id=topic_id)
    else:
        return redirect("view_text_questions_of_topic", topic_id=topic_id)


@login_required   
@staff_member_required(login_url='/login/')
def add_mcq_text_question_view(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)

    # Default: Erster Schritt = Textfrage
    step = request.GET.get("step", "text")
    text_question = None  

    # Schritt 1: Textfrage speichern
    if request.method == "POST" and "text_submit" in request.POST:
     text_form = TextQuestionForm(request.POST)
     if "topic" in text_form.fields:
        text_form.fields.pop("topic")

     if text_form.is_valid():
        text_question = text_form.save(commit=False)
        text_question.topic = topic
        text_question.save()
        messages.success(request, "Textfrage gespeichert. Du kannst weitere hinzufügen oder zu MCQ-Fragen wechseln.")

        # statt sofort zu mcq -> wieder Text-Form zeigen, aber mit text_question für den Button
        text_form = TextQuestionForm()  
        return render(request, "add_mcq_text_question.html", {
            "topic": topic,
            "step": "text",
            "text_form": text_form,
            "mcq_form": None,
            "text_question": text_question,  #  für den Button
        })
        # Initial laden: Schritt 1 oder Schritt 2
    if step == "mcq":
            text_question_id = request.GET.get("text_id")
            if not text_question_id:
                messages.error(request, "Keine Textfrage angegeben. Bitte zuerst eine Textfrage anlegen.")
                return redirect(f"{reverse('add_mcq_text_question_view', args=[topic.topic_id])}?step=text")
            try:
                text_question = TextQuestion.objects.get(pk=text_question_id, topic=topic)
            except TextQuestion.DoesNotExist:
                messages.error(request, "Textfrage nicht gefunden. Bitte erneut anlegen.")
                return redirect(f"{reverse('add_mcq_text_question_view', args=[topic.topic_id])}?step=text")
            if request.method == "POST" and "mcq_submit" in request.POST:
                mcq_form = QuestionForm(request.POST, request.FILES, current_topic=topic)
                if mcq_form.is_valid():
                    question = mcq_form.save(commit=False)
                    question.topic_id = topic
                    question.save()
                    messages.success(request, "MCQ-Frage gespeichert.")
                    # Nach dem Speichern: Entweder weitere MCQ-Frage oder zurück zur Übersicht
                    if "save_continue" in request.POST:
                        return redirect(f"{reverse('add_mcq_text_question_view', args=[topic.topic_id])}?step=mcq&text_id={text_question_id}")
                    else:
                        return redirect("view_mcq_text_questions", topic_id=topic.topic_id)
                else:
                    messages.error(request, "Bitte korrigiere die Fehler im MCQ-Formular.")
            else:
                mcq_form = QuestionForm(current_topic=topic)
            return render(request, "add_mcq_text_question.html", {
                "topic": topic,
                "step": "mcq",
                "text_form": None,
                "mcq_form": mcq_form,
                "text_question": text_question
            })


    # Default: Schritt 1 (Textfrage anzeigen)
    text_form = TextQuestionForm()
    return render(request, "add_mcq_text_question.html", {
        "topic": topic,
        "step": "text",
        "text_form": text_form,
        "mcq_form": None,
    })

@login_required   
@staff_member_required(login_url='/login/')
def view_mcq_text_questions(request, topic_id):
    topic = get_object_or_404(Topics, pk=topic_id)

    # Textfragen dieses Themas laden
    text_questions = TextQuestion.objects.filter(topic=topic).prefetch_related("answers")

    # MCQ-Fragen dieses Themas laden
    mcq_questions = Questions.objects.filter(topic_id=topic)

    return render(
        request,
        "view_mcq_text_questions.html",
        {
            "topic": topic,
            "text_questions": text_questions,
            "mcq_questions": mcq_questions,
        },
    )

@login_required   
@staff_member_required(login_url='/login/')
def view_add_mcq_questions(request, topic_id):
        topic = get_object_or_404(Topics, pk=topic_id)

        if request.method == "POST":
            form = QuestionForm(request.POST, request.FILES)
            # topic_id darf nicht geändert werden, daher Feld deaktivieren
            if "topic_id" in form.fields:
                form.fields["topic_id"].disabled = True
            if form.is_valid():
                question = form.save(commit=False)
                question.topic_id = topic  # Thema bleibt unverändert
                question.save()
                if "save_continue" in request.POST:
                    return redirect("view_add_mcq_questions", topic_id=topic.pk)
                else:
                    return redirect("topic_questions", topic_id=topic.pk)
        else:
            form = QuestionForm()
            if "topic_id" in form.fields:
                form.fields["topic_id"].disabled = True

        return render(request, "add_edit_question.html", {
            "form": form,
            "title": f"Neue MCQ-Frage für: {topic.topic}",
            "topic": topic
        })


# views.py
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect
from django.contrib import messages
from .forms import UserForm, UserProfileForm


@login_required   
def user_profile_view(request):
    user = request.user
    profile, created = UserProfile.objects.get_or_create(user=user)

    # determine template by role
    is_azubi = user.groups.filter(name__iexact="azubis").exists()
    template_name = "user_profile_azubi.html" if is_azubi else "user_profile.html"

    if request.method == "POST":
        user_form = UserForm(request.POST, instance=user)
        profile_form = UserProfileForm(request.POST, request.FILES, instance=profile)

        if user_form.is_valid() and profile_form.is_valid():
            user_form.save()
            profile = profile_form.save(commit=False)
            # Farben extra speichern, falls sie aus JS-Hidden-Feldern kommen
            profile.gradiantcolor1 = request.POST.get('gradiantcolor1', profile.gradiantcolor1)
            profile.gradiantcolor2 = request.POST.get('gradiantcolor2', profile.gradiantcolor2)
            profile.cover_angle = request.POST.get('cover_angle', profile.cover_angle)
            profile.save()

            messages.success(request, "Profil erfolgreich aktualisiert!")
            # Kein redirect – stattdessen Template neu rendern mit aktuellen Daten
        else:
            messages.error(request, "Bitte überprüfe deine Eingaben.")
    else:
        user_form = UserForm(instance=user)
        profile_form = UserProfileForm(instance=profile)

    # Profil nach Speichern neu laden (wichtig)
    profile.refresh_from_db()

    return render(request, template_name, {
        "user_form": user_form,
        "profile_form": profile_form,
        "user_profile": profile,
    })

@user_passes_test(lambda u: u.groups.filter(name__iexact='azubis').exists(), login_url='/')
@login_required
def user_dashboard(request):
    user = request.user

    sessions = (
        TestSession.objects.filter(user=user)
        .select_related("topic__topic_type")
        .prefetch_related("quiz_results__wrong_answers", "text_answers__question")
        .order_by("-started_at")
    )

    results = []

    for s in sessions:
        quiz = s.quiz_results.first()
        text_answers = s.text_answers.all()

        quiz_score = quiz.score if quiz else 0
        quiz_max = quiz.maxsize if quiz else 0

        text_score = sum(a.score for a in text_answers if a.score is not None)
        text_max = sum(a.question.max_score for a in text_answers)

        total_score = quiz_score + text_score
        total_max = quiz_max + text_max
        percent = int((total_score / total_max) * 100) if total_max > 0 else 0

        

        

        results.append({
            "topic": s.topic.topic,
            "type": s.topic.topic_type.name if s.topic.topic_type else "Unbekannt",
            "quiz_score": quiz_score,
            "quiz_max": quiz_max,
            "wrong_answers": quiz.wrong_answers.all() if quiz else [],
            "text_answers": text_answers,
            "total_score": total_score,
            "total_max": total_max,
            "percent": percent,
            "created_at": localtime(s.started_at + timedelta(hours=2)), # wann (Datum/Uhrzeit) anzeigen
        })

    remarks = Remark.objects.filter(user=user).order_by("-created_at")

    return render(request, "azubi_dashboard.html", {
        "results": results,
        "remarks": remarks
    })
